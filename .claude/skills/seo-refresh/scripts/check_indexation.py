"""
check_indexation.py — Indexation and site-health checks.

Two jobs, because ranking analysis alone misses the biggest problems: an
unindexed page cannot rank at all, and a page that 404s is worse than thin.

  1. Parse the Search Console Coverage exports from ~/Downloads:
       *Coverage-*.xlsx            (indexed vs not-indexed trend + issue counts)
       *Coverage-Drilldown-*.xlsx  (the actual URL list for one issue)
  2. Crawl a Hugo build output for broken internal links, baseURL leakage,
     and sitemap sanity.

Usage:
  py .claude/skills/seo-refresh/scripts/check_indexation.py
  py .claude/skills/seo-refresh/scripts/check_indexation.py --build <dir>

To produce a build dir (Hugo is not installed on this machine):
  download hugo_extended_0.160.1_windows-amd64.zip, then
  hugo.exe --gc --minify --destination "$TEMP/bh_verify"
"""

import glob
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")

DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
SITE = "https://biomimicry-hub.com"


def newest(pattern):
    hits = glob.glob(os.path.join(DOWNLOADS, pattern))
    return max(hits, key=os.path.getmtime) if hits else None


def coverage_report():
    try:
        import openpyxl
    except ImportError:
        print("  openpyxl missing — skipping coverage parse")
        return

    path = newest("*Coverage-2*.xlsx") or newest("*Coverage-[0-9]*.xlsx")
    if not path:
        print("  no Coverage export found in Downloads")
        print("  Search Console > Indexing > Pages > Export")
        return

    print("  source: " + os.path.basename(path))
    wb = openpyxl.load_workbook(path, data_only=True)

    if "Chart" in wb.sheetnames:
        rows = [r for r in wb["Chart"].iter_rows(min_row=2, values_only=True)
                if r[1] not in ("", None)]
        if rows:
            first, last = rows[0], rows[-1]
            d_idx = last[2] - first[2]
            print("\n  INDEXATION TREND")
            print("    %s: %d indexed / %d not indexed"
                  % (first[0], first[2], first[1]))
            print("    %s: %d indexed / %d not indexed"
                  % (last[0], last[2], last[1]))
            arrow = "IMPROVING" if d_idx > 0 else ("DECLINING" if d_idx < 0 else "flat")
            print("    indexed %+d  -> %s" % (d_idx, arrow))
            if d_idx < 0:
                print("    ** Indexed pages are falling. This outranks any ranking")
                print("       work: an unindexed page cannot rank at all.")

    for sheet in ("Critical issues", "Non-critical issues"):
        if sheet not in wb.sheetnames:
            continue
        rows = [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r[0]]
        if not rows:
            continue
        print("\n  " + sheet.upper())
        for r in rows:
            print("    %4d pages  %s" % (r[3], r[0]))
            if "404" in str(r[0]):
                print("           -> export the drilldown to see which URLs")
            if "Crawled" in str(r[0]):
                print("           -> thin-content signal; these are deepen candidates")
            if "Alternate page" in str(r[0]):
                print("           -> usually benign, but check it isn't auto-generated")
                print("              taxonomy pages competing with real content")

    drill = newest("*Coverage-Drilldown-*.xlsx")
    if drill:
        wb2 = openpyxl.load_workbook(drill, data_only=True)
        if "Table" in wb2.sheetnames:
            urls = [r[0] for r in wb2["Table"].iter_rows(min_row=2, values_only=True) if r[0]]
            print("\n  DRILLDOWN: %s (%d URLs)" % (os.path.basename(drill), len(urls)))
            # Group by the path prefix after the domain to spot systematic causes.
            groups = {}
            for u in urls:
                p = u.replace(SITE, "").strip("/").split("/")
                groups.setdefault(p[0] if p else "(root)", []).append(u)
            for k, v in sorted(groups.items(), key=lambda x: -len(x[1])):
                print("    %4d  /%s/..." % (len(v), k))
            if len(groups) == 1 or max(len(v) for v in groups.values()) > len(urls) * 0.6:
                print("    ** A single prefix dominates — this is one systematic")
                print("       cause, not scattered rot. Find the config that made it.")


def crawl_build(out):
    print("  build: " + out)
    built = set()
    for f in glob.glob(out + "/**/index.html", recursive=True):
        rel = os.path.relpath(f, out).replace(os.sep, "/").replace("/index.html", "")
        built.add("/" + rel + "/" if rel != "index.html" else "/")
    built.add("/")

    links = {}
    for f in glob.glob(out + "/**/index.html", recursive=True):
        html = open(f, encoding="utf-8", errors="ignore").read()
        src = "/" + os.path.relpath(f, out).replace(os.sep, "/").replace("/index.html", "")
        for m in re.findall(r'href=["\']?(/[^"\'> ]*)', html):
            u = m.split("#")[0].split("?")[0]
            if not u or u in ("/", "//") or u.endswith(
                    (".xml", ".txt", ".svg", ".png", ".jpg", ".css", ".js")):
                continue
            if not u.endswith("/"):
                u += "/"
            links.setdefault(u, set()).add(src)

    print("    built pages: %d   internal link targets: %d" % (len(built), len(links)))
    broken = {u: s for u, s in links.items() if u not in built}
    if broken:
        print("\n    BROKEN INTERNAL LINKS")
        for u, srcs in sorted(broken.items(), key=lambda x: -len(x[1])):
            print("      %s  <- linked from %d page(s)" % (u, len(srcs)))
    else:
        print("    broken internal links: none")

    # baseURL leakage: the GitHub Pages project path must never appear.
    leaked = [f for f in glob.glob(out + "/**/*.html", recursive=True)
              if "/biomimicry-hub/" in open(f, encoding="utf-8", errors="ignore").read()]
    print("    pages containing a /biomimicry-hub/ path prefix: %d%s"
          % (len(leaked), "  ** baseURL regression" if leaked else ""))

    sm = os.path.join(out, "sitemap.xml")
    if os.path.exists(sm):
        locs = re.findall(r"<loc>(.*?)</loc>", open(sm, encoding="utf-8").read())
        bad = [u for u in locs if not u.startswith(SITE + "/")]
        print("    sitemap entries: %d   off-domain: %d" % (len(locs), len(bad)))
        orphan = [u for u in locs if u.replace(SITE, "") not in built]
        if orphan:
            print("    ** %d sitemap URLs have no built page" % len(orphan))


def main():
    args = sys.argv[1:]
    print("=" * 64)
    print("COVERAGE / INDEXATION")
    print("=" * 64)
    coverage_report()

    build = args[args.index("--build") + 1] if "--build" in args else None
    if not build:
        guess = os.path.join(os.environ.get("TEMP", "/tmp"), "bh_verify")
        build = guess if os.path.isdir(guess) else None
    print("\n" + "=" * 64)
    print("BUILD / LINK HEALTH")
    print("=" * 64)
    if build and os.path.isdir(build):
        crawl_build(build)
    else:
        print("  no build dir — pass --build <dir> after running Hugo")


if __name__ == "__main__":
    main()
